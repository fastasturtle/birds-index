#!/usr/bin/env python3
"""Build data/species.json: a trilingual (ru / en / Latin) index of world bird species.

Backbone: AviList v2025b (species only). Russian names are taken from a cascade
of sources, matched by exact (normalised) scientific name:

    1. IOC World Bird List Multilingual 15.2   -> ru_source "ioc"
    2. eBird taxonomy API, locale=ru            -> ru_source "ebird"
    3. Wikidata P1843 (taxon common name, ru)   -> ru_source "wikidata"
    4. data/overrides.csv (applied last)        -> ru_source "manual"

Then the Russian name is normalised: synonyms ("Хохлатый фазан, чир",
"Гуменник (лесной гуменник)") go from `ru` to the list `ru_alt`, and index
fields are computed: `ru_head` (leading noun), `ru_book` (book spelling with a
lower-case first letter), `en_head` (last word of the English name).

It also writes data/taxa.json: orders and families (in systematic order) with
Russian names from the ZMMU passerine spreadsheet (family header rows) ->
"zmmu", Wikidata (P1843 / ru label of the order or family item) -> "wikidata",
and data/overrides.csv (sci = order or family name, field = ru) -> "manual".

Usage:
    uv run scripts/build.py            # download whatever is missing, then build
    uv run scripts/build.py --offline  # use only files already on disk
    uv run scripts/build.py --fetch    # force re-download of every source
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import sys
import time
from collections import Counter
from pathlib import Path

import openpyxl
import requests

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
RAW = DATA / "raw"
CACHE = DATA / "cache"

AVILIST_PATH = RAW / "AviList-v2025b-extended.xlsx"
IOC_PATH = RAW / "IOC-Multiling-15.2.xlsx"
EBIRD_PATH = CACHE / "ebird-taxonomy-ru.json"
WIKIDATA_PATH = CACHE / "wikidata-ru-p1843.json"
WIKIDATA_TAXA_PATH = CACHE / "wikidata-taxa-ru.json"
ZMMU_PASSERINE_PATH = RAW / "Birds_rus_names_passerine.xlsx"
OVERRIDES_PATH = DATA / "overrides.csv"
OUT_JSON = DATA / "species.json"
OUT_DOCS_JSON = ROOT / "docs" / "data" / "species.json"
OUT_TAXA_JSON = DATA / "taxa.json"
OUT_DOCS_TAXA_JSON = ROOT / "docs" / "data" / "taxa.json"
OUT_REPORT = DATA / "report.md"

AVILIST_URL = "https://www.avilist.org/wp-content/uploads/2026/06/AviList-v2025b-10Jun2026-extended.xlsx"
IOC_URL = "https://worldbirdnames.org/Multiling%20IOC%2015.2.xlsx"
EBIRD_URL = "https://api.ebird.org/v2/ref/taxonomy/ebird?fmt=json&locale=ru&cat=species"
WIKIDATA_SPARQL = "https://query.wikidata.org/sparql"
ZMMU_PASSERINE_URL = "https://zmmu.msu.ru/files/publications/Birds_rus_names_passerine.xlsx"

USER_AGENT = "birds-index-build/0.1 (https://github.com/fastasturtle/birds-index)"
CYRILLIC = re.compile(r"[Ѐ-ӿ]")

EXTINCT_MAP = {
    "extinct": "extinct",
    "(extinct)": "extinct",
    "extinct in the wild": "extinct_in_wild",
    "possibly extinct": "possibly_extinct",
    "(possibly extinct)": "possibly_extinct",
    "possibly extinct in the wild": "possibly_extinct",
    "previously extinct in the wild, but reintroduced": None,
}

FIELDS = [
    "id", "seq", "order", "family", "family_en", "sci", "authority", "auth", "en",
    "ru", "ru_alt", "ru_head", "ru_head_source", "ru_book", "en_head",
    "ru_source", "ru_status", "ioc_sci", "extinct", "iucn",
]
# Fields computed by normalise_names / add_index_fields; only ru_alt and
# ru_head may be given in overrides.csv (handled specially).
DERIVED_FIELDS = {"ru_head_source", "ru_book", "en_head"}


# --------------------------------------------------------------------------- helpers

def sci_key(name: str | None) -> str:
    """Normalise a scientific name for matching: lower case, single spaces."""
    return " ".join(str(name or "").split()).lower()


def clean(value) -> str | None:
    """Strip and collapse whitespace; empty -> None."""
    if value is None:
        return None
    text = " ".join(str(value).split())
    return text or None


def log(msg: str) -> None:
    print(msg, file=sys.stderr)


AUTH_YEAR = re.compile(r"(.*?),\s*(\d{4}[a-z]?)")
AUTH_SUFFIX = re.compile(r"\b(Jr|Sr)\.?$")
AUTH_TRAILING_INITIALS = re.compile(r"\s+[A-Z]{1,3}$")


def normalise_authority(raw: str | None) -> str | None:
    """AviList authority -> conventional zoological form without initials.

    "Linnaeus, C, 1758" -> "Linnaeus, 1758"; "(Latham, J, 1790)" -> "(Latham, 1790)";
    two authors "A & B, year"; three or more "A, B & C, year". Authors are
    separated by ";" and written "Surname, Initials" (Jr/Sr kept as a suffix);
    an author without a comma is kept, minus trailing initials ("Zhao C" ->
    "Zhao"), so full names such as "Liu Yang" or "Suparno" survive. Strings
    that do not end in a year are returned unchanged.
    """
    if raw is None:
        return None
    text = raw.strip()
    paren = text.startswith("(") and text.endswith(")")
    if paren:
        text = text[1:-1].strip()
    m = AUTH_YEAR.fullmatch(text)
    if not m:
        return raw
    names = []
    for part in m.group(1).split(";"):
        part = part.strip()
        if "," in part:
            surname, initials = (x.strip() for x in part.rsplit(",", 1))
            suffix = AUTH_SUFFIX.search(initials)
            if suffix:
                surname = f"{surname} {suffix.group(1)}"
        else:
            surname = AUTH_TRAILING_INITIALS.sub("", part)
        names.append(surname)
    if len(names) == 1:
        authors = names[0]
    else:
        authors = ", ".join(names[:-1]) + " & " + names[-1]
    out = f"{authors}, {m.group(2)}"
    return f"({out})" if paren else out


def _test_normalise_authority() -> None:
    cases = {
        "Linnaeus, C, 1758": "Linnaeus, 1758",
        "(Latham, J, 1790)": "(Latham, 1790)",
        "Sclater, PL; von Hochstetter, F, 1861": "Sclater & von Hochstetter, 1861",
        "Tennyson, AJD; Palma, RL; Robertson, HA; Worthy, TH; Gill, BJ, 2003":
            "Tennyson, Palma, Robertson, Worthy & Gill, 2003",
        "(d'Orbigny, ACVMD; de Lafresnaye, NFAA, 1837)": "(d'Orbigny & de Lafresnaye, 1837)",
        "Phelps, WH Sr; Phelps, WH Jr, 1946": "Phelps Sr & Phelps Jr, 1946",
        "Zhou F; Jiang A, 2008": "Zhou & Jiang, 2008",
        "Indrawan, M; Rasmussen, PC; Sunarto, 2008": "Indrawan, Rasmussen & Sunarto, 2008",
        "(Oustalet, J-FÉ, 1880)": "(Oustalet, 1880)",
        "Mahood, SP; Hong Chamnan; Poole, CM, 2013": "Mahood, Hong Chamnan & Poole, 2013",
    }
    for raw, want in cases.items():
        got = normalise_authority(raw)
        assert got == want, f"normalise_authority({raw!r}) = {got!r}, want {want!r}"
    assert normalise_authority(None) is None


LEFTOVER_INITIALS = re.compile(r", [A-Z]{1,3}[,;)]|, [A-Z]{1,3}$")


def download(url: str, dest: Path, timeout: int = 300) -> None:
    """Download url to dest atomically (via a .part file)."""
    log(f"  downloading {url}")
    dest.parent.mkdir(parents=True, exist_ok=True)
    tmp = dest.with_suffix(dest.suffix + ".part")
    with requests.get(url, stream=True, timeout=timeout,
                      headers={"User-Agent": USER_AGENT}) as r:
        r.raise_for_status()
        with open(tmp, "wb") as f:
            for chunk in r.iter_content(1 << 16):
                f.write(chunk)
    tmp.replace(dest)


def ensure_file(path: Path, url: str, mode: str, required: bool) -> bool:
    """Make sure `path` exists according to mode ('offline' | 'default' | 'fetch').

    Returns True if the file is available. Raises if a required file is missing.
    """
    need = mode == "fetch" or not path.exists()
    if need and mode != "offline":
        try:
            download(url, path)
        except Exception as exc:  # noqa: BLE001 - report and continue
            log(f"  WARNING: download failed for {url}: {exc}")
    if path.exists():
        return True
    if required:
        raise SystemExit(f"Required file missing: {path} (run without --offline to download it)")
    return False


def xlsx_rows(path: Path, sheet: str | None = None):
    """Yield (header, row) dicts from the first (or named) worksheet."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    for row in rows:
        yield dict(zip(header, row))
    wb.close()


# --------------------------------------------------------------------------- sources

def load_avilist(path: Path) -> list[dict]:
    """Return AviList species rows as base output records (in sequence order)."""
    records = []
    for row in xlsx_rows(path):
        if clean(row.get("Taxon_rank")) != "species":
            continue  # also skips the second, human-readable header row
        status = clean(row.get("Extinct_or_possibly_extinct"))
        if status is not None and status.lower() not in EXTINCT_MAP:
            raise SystemExit(f"Unknown extinction status {status!r} for {row.get('Scientific_name')}")
        records.append({
            "id": clean(row.get("AvibaseID")),
            "seq": int(row["Sequence"]),
            "order": clean(row.get("Order")),
            "family": clean(row.get("Family")),
            "family_en": clean(row.get("Family_English_name")),
            "sci": clean(row.get("Scientific_name")),
            "authority": clean(row.get("Authority")),
            "auth": normalise_authority(clean(row.get("Authority"))),
            "en": clean(row.get("English_name_AviList")),
            "ru": None,
            "ru_alt": [],
            "ru_head": None,
            "ru_head_source": None,
            "ru_book": None,
            "en_head": None,
            "ru_source": None,
            "ru_status": "missing",
            "ioc_sci": None,  # exact matching only in this version
            "extinct": EXTINCT_MAP[status.lower()] if status else None,
            "iucn": clean(row.get("IUCN_Red_List_Category")),
        })
    records.sort(key=lambda r: r["seq"])
    return records


def load_ioc(path: Path) -> dict[str, dict]:
    """Return {sci_key: {"sci", "en", "ru"}} from the IOC multilingual list."""
    out = {}
    for row in xlsx_rows(path, sheet="List"):
        sci = clean(row.get("IOC_15.2"))
        if not sci:
            continue
        out[sci_key(sci)] = {"sci": sci, "en": clean(row.get("English")),
                             "ru": clean(row.get("Russian"))}
    return out


def load_ebird(path: Path) -> dict[str, str]:
    """Return {sci_key: Russian name} from the eBird taxonomy (locale=ru).

    eBird falls back to English when no Russian name exists, so names without
    any Cyrillic are treated as missing.
    """
    data = json.loads(path.read_text(encoding="utf-8"))
    out = {}
    for item in data:
        if item.get("category") != "species":
            continue
        name = clean(item.get("comName"))
        if name and CYRILLIC.search(name):
            out[sci_key(item.get("sciName"))] = name
    return out


def sparql(query: str, timeout: int = 180) -> list[dict]:
    r = requests.get(WIKIDATA_SPARQL, params={"query": query, "format": "json"},
                     headers={"User-Agent": USER_AGENT,
                              "Accept": "application/sparql-results+json"},
                     timeout=timeout)
    r.raise_for_status()
    return r.json()["results"]["bindings"]


WIKIDATA_QUERY = """
SELECT ?item ?sci ?ru WHERE {{
  ?item wdt:P105 wd:Q7432 ;
        wdt:P225 ?sci ;
        wdt:P1843 ?ru .
  FILTER(LANG(?ru) = "ru")
  {extra}
}}
"""

WIKIDATA_LABELS = """
SELECT ?item ?label WHERE {{
  VALUES ?item {{ {items} }}
  ?item rdfs:label ?label . FILTER(LANG(?label) = "ru")
}}
"""


def sparql_retry(query: str, attempts: int = 3) -> list[dict]:
    for attempt in range(attempts):
        try:
            return sparql(query)
        except Exception as exc:  # noqa: BLE001
            if attempt == attempts - 1:
                raise
            log(f"  SPARQL request failed ({exc.__class__.__name__}); retrying")
            time.sleep(10 * (attempt + 1))
    raise AssertionError("unreachable")


def fetch_wikidata(dest: Path) -> None:
    """Fetch all species items with a Russian P1843 into a JSON cache.

    Not restricted to birds: the result set is small (~11k rows) and matching
    against AviList by scientific name restricts it anyway. Tries one query
    first; if it fails, splits the query by the first letter of the taxon name.
    Russian labels are then fetched only for taxa with several ru names (used
    to pick one); a failure there is not fatal.
    """
    log("  querying Wikidata SPARQL")
    try:
        bindings = sparql_retry(WIKIDATA_QUERY.format(extra=""), attempts=2)
    except Exception as exc:  # noqa: BLE001
        log(f"  single query failed ({exc.__class__.__name__}); splitting by initial letter")
        bindings = []
        for letter in "ABCDEFGHIJKLMNOPQRSTUVWXYZ":
            bindings += sparql_retry(WIKIDATA_QUERY.format(
                extra=f'FILTER(STRSTARTS(UCASE(STR(?sci)), "{letter}"))'))
        # names not starting with A-Z (rare; e.g. leading quote); best effort
        try:
            bindings += sparql_retry(WIKIDATA_QUERY.format(
                extra='FILTER(!REGEX(STR(?sci), "^[A-Za-z]"))'))
        except Exception as exc2:  # noqa: BLE001
            log(f"  non-A-Z query failed ({exc2.__class__.__name__}); ignoring")
    rows = [{
        "item": b["item"]["value"].rsplit("/", 1)[-1],
        "sci": b["sci"]["value"],
        "ru": b["ru"]["value"],
        "label": None,
    } for b in bindings]

    names_per_item = Counter(r["item"] for r in rows)
    multi = sorted(i for i, n in names_per_item.items() if n > 1)
    labels = {}
    try:
        for i in range(0, len(multi), 200):
            batch = " ".join(f"wd:{q}" for q in multi[i:i + 200])
            for b in sparql_retry(WIKIDATA_LABELS.format(items=batch)):
                labels[b["item"]["value"].rsplit("/", 1)[-1]] = b["label"]["value"]
    except Exception as exc:  # noqa: BLE001
        log(f"  WARNING: label query failed ({exc.__class__.__name__}); continuing without labels")
    for r in rows:
        r["label"] = labels.get(r["item"])

    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")
    log(f"  cached {len(rows)} Wikidata rows ({len(multi)} items with several ru names)")


def load_wikidata(path: Path) -> dict[str, str]:
    """Return {sci_key: Russian name}. If several ru names exist for a taxon,
    prefer the one equal to the item's Russian label, else the first sorted."""
    rows = json.loads(path.read_text(encoding="utf-8"))
    by_sci: dict[str, list[tuple[str, str | None]]] = {}
    for row in rows:
        name = clean(row.get("ru"))
        if name:
            by_sci.setdefault(sci_key(row["sci"]), []).append((name, clean(row.get("label"))))
    out = {}
    for key, cands in by_sci.items():
        labels = {lab for _, lab in cands if lab}
        names = sorted({n for n, _ in cands})
        preferred = [n for n in names if n in labels]
        out[key] = (preferred or names)[0]
    return out


# ---- order / family names

ZMMU_HEADER = re.compile(r"([A-Z][a-z]+idae)\b\s*[-–—]?\s*(.*)")
# Stressed vowels in the ZMMU file are Latin look-alikes: "Стрелкú" (ú = и́),
# "Яхý" (ý = у́).
STRESS_MARKS = str.maketrans({"á": "а", "é": "е", "ó": "о", "ú": "и", "ý": "у", "í": "и",
                              "\u0301": None})


def normalise_taxon_ru(name: str | None) -> str | None:
    """Strip whitespace, drop stress marks, capitalise the first letter."""
    name = clean((name or "").translate(STRESS_MARKS))
    if not name:
        return None
    return name[0].upper() + name[1:]


def load_zmmu_families(path: Path, species_family: dict[str, str],
                       families: set[str]) -> tuple[dict[str, str], list[str]]:
    """Return ({family_key: Russian name}, warnings) from the ZMMU passerine list.

    Family header rows look like "Pittidae - Питтовые" (alternative names in
    parentheses are dropped). Each header is checked against the AviList
    families of the species listed under it: if the header's Latin name is not
    an AviList family, or none of those species belongs to it (slips such as
    a misspelt or wrongly repeated header), the name is given to the majority
    family of those species instead, but only if that family has no header of
    its own.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    blocks: list[list] = []  # [latin, russian, Counter of AviList families]
    for row in ws.iter_rows(values_only=True):
        a = clean(row[0]) if row else None
        rest = [clean(v) for v in row[1:4]] if row else []
        if not a:
            continue
        m = ZMMU_HEADER.fullmatch(a)
        if m and not any(rest):
            ru = re.sub(r"\([^)]*\)", "", m.group(2))
            blocks.append([m.group(1), normalise_taxon_ru(ru), Counter()])
        elif blocks and sci_key(a) in species_family:
            blocks[-1][2][species_family[sci_key(a)]] += 1
    wb.close()

    out: dict[str, str] = {}
    warnings: list[str] = []
    deferred = []
    for latin, ru, fams in blocks:
        if not ru:
            continue
        majority = fams.most_common(1)[0][0] if fams else None
        if latin in families and (not fams or fams[latin] > 0):
            out.setdefault(latin.lower(), ru)
        else:
            deferred.append((latin, ru, majority))
    for latin, ru, majority in deferred:
        if majority and majority.lower() not in out:
            out[majority.lower()] = ru
            warnings.append(f"ZMMU header {latin!r} ({ru}) lists {majority} species; "
                            f"used for {majority}")
        else:
            warnings.append(f"ZMMU header {latin!r} ({ru}) not used "
                            f"(majority AviList family of its species: {majority or 'none'})")
    return out, warnings


WIKIDATA_TAXA_QUERY = """
SELECT ?item ?sci ?rank ?ru ?label ?aves WHERE {{
  VALUES ?sci {{ {names} }}
  VALUES ?rank {{ wd:Q36602 wd:Q35409 }}
  ?item wdt:P225 ?sci ;
        wdt:P105 ?rank .
  OPTIONAL {{ ?item wdt:P1843 ?ru . FILTER(LANG(?ru) = "ru") }}
  OPTIONAL {{ ?item rdfs:label ?label . FILTER(LANG(?label) = "ru") }}
  BIND(EXISTS {{ ?item wdt:P171+ wd:Q5113 }} AS ?aves)
}}
"""


def fetch_wikidata_taxa(dest: Path, names: list[str]) -> None:
    """One SPARQL query (VALUES list) for the order/family items; cache raw rows."""
    log(f"  querying Wikidata SPARQL for {len(names)} order/family names")
    values = " ".join(json.dumps(n) for n in names)
    bindings = sparql_retry(WIKIDATA_TAXA_QUERY.format(names=values), attempts=2)
    rows = [{
        "item": b["item"]["value"].rsplit("/", 1)[-1],
        "sci": b["sci"]["value"],
        "rank": "order" if b["rank"]["value"].endswith("/Q36602") else "family",
        "ru": b.get("ru", {}).get("value"),
        "label": b.get("label", {}).get("value"),
        "aves": b.get("aves", {}).get("value") == "true",
    } for b in bindings]
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(json.dumps({"queried": sorted(names), "rows": rows},
                               ensure_ascii=False, indent=1), encoding="utf-8")
    log(f"  cached {len(rows)} Wikidata order/family rows")


def load_wikidata_taxa(path: Path) -> dict[tuple[str, str], str]:
    """Return {(rank, sci_key): Russian name}.

    Items placed under Aves win over homonyms elsewhere. P1843 (ru) is
    preferred, choosing the value equal to the ru label if there are several
    (else the first sorted); otherwise the ru label if it contains Cyrillic.
    """
    rows = json.loads(path.read_text(encoding="utf-8"))["rows"]
    items: dict[tuple[str, str], dict[str, dict]] = {}
    for r in rows:
        it = items.setdefault((r["rank"], sci_key(r["sci"])), {}).setdefault(
            r["item"], {"ru": set(), "label": None, "aves": False})
        if r.get("ru"):
            it["ru"].add(normalise_taxon_ru(r["ru"]))
        if r.get("label"):
            it["label"] = normalise_taxon_ru(r["label"])
        it["aves"] = it["aves"] or bool(r.get("aves"))
    out = {}
    for key, cands in items.items():
        pool = [c for c in cands.values() if c["aves"]] or list(cands.values())
        for c in sorted(pool, key=lambda c: (not c["ru"], not c["label"])):
            names = sorted(c["ru"])
            if names:
                out[key] = c["label"] if c["label"] in names else names[0]
                break
            if c["label"] and CYRILLIC.search(c["label"]):
                out[key] = c["label"]
                break
    return out


def build_taxa(records) -> dict[str, list[dict]]:
    """Orders and families with species counts and first sequence number."""
    orders: dict[str, dict] = {}
    families: dict[str, dict] = {}
    for r in records:
        o = orders.setdefault(r["order"], {"sci": r["order"], "ru": None, "ru_source": None,
                                           "seq": r["seq"], "count": 0})
        o["seq"] = min(o["seq"], r["seq"])
        o["count"] += 1
        f = families.setdefault(r["family"], {"sci": r["family"], "en": r["family_en"],
                                              "ru": None, "ru_source": None,
                                              "order": r["order"], "seq": r["seq"],
                                              "count": 0})
        f["seq"] = min(f["seq"], r["seq"])
        f["count"] += 1
    return {"orders": sorted(orders.values(), key=lambda t: t["seq"]),
            "families": sorted(families.values(), key=lambda t: t["seq"])}


def apply_taxa_names(taxa, zmmu: dict[str, str], wikidata_taxa: dict) -> None:
    for rank, items in (("order", taxa["orders"]), ("family", taxa["families"])):
        for t in items:
            key = sci_key(t["sci"])
            for source, name in (("zmmu", zmmu.get(key) if rank == "family" else None),
                                 ("wikidata", wikidata_taxa.get((rank, key)))):
                if name:
                    t["ru"], t["ru_source"] = name, source
                    break


def apply_taxa_overrides(taxa, overrides) -> list[str]:
    by_key = {sci_key(t["sci"]): t for t in taxa["orders"] + taxa["families"]}
    warnings = []
    for ov in overrides:
        t = by_key[sci_key(ov["sci"])]
        field = (ov.get("field") or "").strip()
        if field != "ru":
            warnings.append(f"override with unsupported field {field!r} for taxon {ov['sci']!r}")
            continue
        t["ru"] = normalise_taxon_ru(ov.get("value"))
        t["ru_source"] = "manual" if t["ru"] else None
    return warnings


def load_overrides(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with open(path, newline="", encoding="utf-8") as f:
        return [row for row in csv.DictReader(f) if (row.get("sci") or "").strip()]


# --------------------------------------------------------------------------- build

def apply_cascade(records, ioc, ebird, wikidata):
    for rec in records:
        key = sci_key(rec["sci"])
        for source, table in (("ioc", ioc), ("ebird", ebird), ("wikidata", wikidata)):
            if source == "ioc":
                name = (table.get(key) or {}).get("ru")
            else:
                name = table.get(key)
            if name:
                rec["ru"], rec["ru_source"] = name, source
                break


def apply_overrides(records, overrides, taxon_keys=frozenset()) -> list[str]:
    """Apply manual species overrides in file order; return a list of warnings.

    Rows whose `sci` is an order or family name (taxon_keys) are skipped here;
    apply_taxa_overrides handles them.
    """
    by_key = {sci_key(r["sci"]): r for r in records}
    warnings = []
    for ov in overrides:
        if sci_key(ov["sci"]) in taxon_keys and sci_key(ov["sci"]) not in by_key:
            continue
        rec = by_key.get(sci_key(ov["sci"]))
        field = (ov.get("field") or "").strip()
        value = clean(ov.get("value"))
        if rec is None:
            warnings.append(f"override for unknown species {ov['sci']!r}")
            continue
        if field not in FIELDS or field in ("ru_source", "ru_status") or field in DERIVED_FIELDS:
            warnings.append(f"override with unsupported field {field!r} for {ov['sci']!r}")
            continue
        if field == "ru_alt":
            # "a; b" -> accumulated on top of the automatic alternatives
            rec.setdefault("_alt_manual", []).extend(
                cap_first(x) for x in (clean(v) for v in (value or "").split(";")) if x)
            continue
        if field == "ru_head":
            rec["_head_manual"] = value
            continue
        if field == "seq":
            value = int(value) if value else rec["seq"]
        rec[field] = value
        if field == "ru":
            rec["ru_source"] = "manual" if value else None
    return warnings


# --------------------------------------------------------------------------- name normalisation

def cap_first(text: str) -> str:
    return text[:1].upper() + text[1:] if text else text


LATIN = re.compile(r"[A-Za-z]")
MIXED_WORD = re.compile(r"\S+")


def mixed_script_words(name: str | None) -> list[str]:
    """Words mixing Cyrillic and Latin letters (look-alike typos: "Cкоморох")."""
    return [w for w in MIXED_WORD.findall(name or "")
            if CYRILLIC.search(w) and LATIN.search(w)]


# "Гуменник (лесной гуменник)": the parenthesised part is taken as an
# alternative name only if it looks like a name, not like a comment.
PAREN = re.compile(r"^(?P<main>[^()]+?)\s*\((?P<alt>[^()]+)\)\s*$")
PAREN_NOT_NAME = re.compile(r"[A-Za-z0-9=?.:;]|\b(?:вид|подвид|подвиды|форма|группа|включая|"
                            r"кроме|см|ранее|прежде|s\.l|возможно)\b", re.IGNORECASE)
ALT_SPLIT = re.compile(r"\s*,\s*|\s+/\s+|\s+или\s+")


def split_ru(name: str) -> tuple[str, list[str], str | None]:
    """Split a Russian name with synonyms into (main, alternatives, problem).

    "Хохлатый фазан, чир" -> ("Хохлатый фазан", ["Чир"]); separators: comma,
    " / ", " или ". "Гуменник (лесной гуменник)" -> ("Гуменник", ["Лесной
    гуменник"]). A parenthesis that does not look like a name (Latin, digits,
    punctuation, words such as "подвид", more than four words) is left as is
    and reported via `problem`.
    """
    alts: list[str] = []
    m = PAREN.match(name)
    if m:
        inner = clean(m.group("alt"))
        if PAREN_NOT_NAME.search(inner) or len(inner.split()) > 4:
            return name, [], f"parenthesis kept: {name!r}"
        name, alts = clean(m.group("main")), [inner]
    elif "(" in name or ")" in name:
        return name, [], f"parenthesis kept: {name!r}"
    parts = [p for p in (clean(x) for x in ALT_SPLIT.split(name)) if p]
    return cap_first(parts[0]), [cap_first(p) for p in parts[1:] + alts], None


# Adjective endings, incl. participles in -ся ("кланяющаяся") and possessive
# adjectives in -ья/-ье/-ьи/-ий ("сорочья", "лисья").
ADJ_END = re.compile(r"(?:ый|ий|ой|ая|яя|ое|ее|ые|ие|ья|ье|ьи|ьё)(?:ся)?$")
# Possessive adjectives: "Стеллеров", "Бэрова", "дарвинов", "буйволов".
POSSESSIVE_END = re.compile(r"(?:ов|ова|ово|ев|ева|ево|ёв|ёва|ин|ина|ино)$")
# Possessives formed from common nouns: lower case in book spelling.
COMMON_POSSESSIVES = {"крокодилов", "фараонов", "буйволов"}


def ru_head(name: str | None) -> tuple[str | None, bool]:
    """Leading word ("nest") of a Russian species name, and a doubt flag.

    Words before the first noun are dropped if they are adjectives (endings
    -ый/-ий/-ой/-ая/-яя/-ое/-ее/-ые/-ие, -ья/-ье, participles in -ся) or
    possessive adjectives (-ов(а)/-ев(а)/-ин(а)) followed by a lower-case word
    ("Стеллеров морской орёл", "буйволов ткач"; but "Пингвин Адели"). The first
    remaining word is the head; everything after it is a genitive or an
    apposition ("гусь Росса", "фазан микадо", "конёк пуны"). Hyphenated words
    count as one ("лебедь-шипун"). Flagged as doubtful: every word looks like
    an adjective (the last word is taken), a lower-case word follows the head
    (apposition, or a typo in an adjective), or a capitalised head after
    other words. The head is returned with a lower-case first letter.
    """
    if not name:
        return None, False
    words = name.split()
    head, doubtful = words[-1], True
    for i, w in enumerate(words):
        base = w.lower()
        # an adjective is always followed by a lower-case word; before a
        # capitalised genitive the word is a noun ("Козодой Пригожина")
        if i + 1 < len(words) and words[i + 1][:1].islower() and (
                ADJ_END.search(base) or POSSESSIVE_END.search(base)):
            continue
        head = w
        # doubtful: a capitalised head after other words, or a lower-case
        # word after the head (apposition such as "конёк пуны", or a typo)
        doubtful = (i > 0 and w[:1].isupper()) or (
            i + 1 < len(words) and words[i + 1][:1].islower())
        break
    return head[:1].lower() + head[1:], doubtful


def en_head(name: str | None) -> str | None:
    """Last word of the English name (a hyphenated word is kept whole)."""
    if not name:
        return None
    return name.split()[-1]


def ru_book(name: str | None, en: str | None) -> tuple[str | None, bool]:
    """Russian name in book spelling: first letter lower case, unless the first
    word is a possessive adjective formed from a proper name ("Стеллерова
    сойка", "Магелланов гусь"): a capitalised word in -ов(а)/-ев(а)/-ин(а)
    followed by a lower-case word ("Пингвин Адели" is not one). The English
    "'s" ("Steller's Jay") confirms it; without it the capital is kept unless
    the word is a known possessive of a common noun (COMMON_POSSESSIVES:
    "Крокодилов сторож", "Фараонов филин"). Other words are unchanged
    ("гусь Росса"). Returns (name, capital kept)."""
    if not name:
        return None, False
    words = name.split()
    first = words[0].lower()
    if (len(words) > 1 and words[0][:1].isupper() and words[1][:1].islower()
            and POSSESSIVE_END.search(first) and not ADJ_END.search(first)
            and ("'s" in (en or "") or first not in COMMON_POSSESSIVES)):
        return name, True
    return name[:1].lower() + name[1:], False


def _test_names() -> None:
    assert split_ru("Хохлатый фазан, чир") == ("Хохлатый фазан", ["Чир"], None)
    assert split_ru("Гуменник (лесной гуменник)") == ("Гуменник", ["Лесной гуменник"], None)
    assert split_ru("Синьга") == ("Синьга", [], None)
    assert split_ru("Птица (подвид X)")[2]
    for name, want in {"Белолобый гусь": "гусь", "Лебедь-шипун": "лебедь-шипун",
                       "Гусь Росса": "гусь", "Малая канадская казарка": "казарка",
                       "Эму": "эму", "Пеганка-раджа": "пеганка-раджа",
                       "Стеллеров морской орёл": "орёл"}.items():
        got = ru_head(name)
        assert got == (want, False), f"ru_head({name!r}) = {got!r}, want {want!r}"
    assert en_head("Eastern Whip-poor-will") == "Whip-poor-will"
    assert en_head("Emu") == "Emu"
    assert en_head("Florida Scrub-Jay") == "Scrub-Jay"
    assert ru_book("Белолобый гусь", "Greater White-fronted Goose") == ("белолобый гусь", False)
    assert ru_book("Гусь Росса", "Ross's Goose") == ("гусь Росса", False)
    assert ru_book("Стеллеров морской орёл", "Steller's Sea Eagle") == ("Стеллеров морской орёл", True)
    assert ru_book("Розовый пеликан", "Great White Pelican")[1] is False
    assert ru_book("Магелланов гусь", "Upland Goose") == ("Магелланов гусь", True)
    assert ru_book("Крокодилов сторож", "Egyptian Plover") == ("крокодилов сторож", False)
    assert ru_book("Пингвин Адели", "Adelie Penguin") == ("пингвин Адели", False)
    for name, want in {"Бурая кланяющаяся крачка": "крачка",
                       "Сорочья кукушка": "кукушка", "Красноклювый буйволов ткач": "ткач",
                       "Пингвин Адели": "пингвин", "Козодой Пригожина": "козодой"}.items():
        assert ru_head(name) == (want, False), (name, ru_head(name))
    assert ru_head("Фазан микадо") == ("фазан", True)


def normalise_names(records) -> dict[str, list]:
    """Split synonyms out of `ru` into `ru_alt`, merge manual ru_alt, compute
    ru_head / ru_book / en_head. Returns lists for the build report."""
    log_ = {"split": [], "paren_kept": [], "mixed": [], "head_doubtful": [],
            "book_capital": []}
    for rec in records:
        alts: list[str] = []
        if rec["ru"]:
            main, auto, problem = split_ru(rec["ru"])
            if problem:
                log_["paren_kept"].append((rec["sci"], rec["ru"]))
            elif auto:
                log_["split"].append((rec["sci"], rec["ru"], main, "; ".join(auto),
                                      rec["ru_source"]))
            rec["ru"] = main
            alts = auto
        for a in alts + rec.pop("_alt_manual", []):
            if a not in rec["ru_alt"] and a != rec["ru"]:
                rec["ru_alt"].append(a)
        for name in [rec["ru"], rec["en"], *rec["ru_alt"]]:
            if mixed_script_words(name):
                log_["mixed"].append((rec["sci"], name))
        manual = rec.pop("_head_manual", None)
        if manual and rec["ru"]:
            rec["ru_head"], rec["ru_head_source"] = manual, "manual"
        elif rec["ru"]:
            rec["ru_head"], doubtful = ru_head(rec["ru"])
            rec["ru_head_source"] = "auto"
            if doubtful:
                log_["head_doubtful"].append((rec["sci"], rec["ru"], rec["ru_head"]))
        else:
            rec["ru_head"] = rec["ru_head_source"] = None
        rec["ru_book"], kept = ru_book(rec["ru"], rec["en"])
        if kept:
            log_["book_capital"].append((rec["sci"], rec["ru"], rec["en"]))
        rec["en_head"] = en_head(rec["en"])
    return log_


def finalise(records):
    for rec in records:
        if rec["ru"] is None:
            rec["ru_source"], rec["ru_status"] = None, "missing"
        elif rec["ru_source"] in ("ioc", "manual"):
            rec["ru_status"] = "ok"
        else:
            rec["ru_status"] = "ok_fallback"


def write_json(records, taxa):
    payload = json.dumps([{k: r[k] for k in FIELDS} for r in records],
                         ensure_ascii=False, separators=(",", ":"))
    OUT_JSON.write_text(payload, encoding="utf-8")
    OUT_TAXA_JSON.write_text(json.dumps(taxa, ensure_ascii=False, separators=(",", ":")),
                             encoding="utf-8")
    OUT_DOCS_JSON.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(OUT_JSON, OUT_DOCS_JSON)
    shutil.copyfile(OUT_TAXA_JSON, OUT_DOCS_TAXA_JSON)


def md_table(rows, headers) -> list[str]:
    def esc(v):
        return "" if v is None else str(v).replace("|", "\\|")
    out = ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    out += ["| " + " | ".join(esc(v) for v in row) + " |" for row in rows]
    return out


def duplicate_ru(records) -> list[tuple[str, list[dict]]]:
    """Russian names (case-insensitive) shared by several species."""
    by_name: dict[str, list[dict]] = {}
    for r in records:
        if r["ru"]:
            by_name.setdefault(r["ru"].casefold(), []).append(r)
    return [(recs[0]["ru"], recs) for recs in by_name.values() if len(recs) > 1]


def alt_clashes(records) -> list[tuple[str, str, str]]:
    """ru_alt entries equal to the main Russian name of another species."""
    main = {r["ru"].casefold(): r["sci"] for r in records if r["ru"]}
    return [(r["sci"], a, main[a.casefold()]) for r in records for a in r["ru_alt"]
            if a.casefold() in main and main[a.casefold()] != r["sci"]]


def write_report(records, ioc, ebird, wikidata, status, warnings, taxa, taxa_warnings,
                 names_log) -> dict:
    avi_keys = {sci_key(r["sci"]) for r in records}
    total = len(records)
    by_source = Counter(r["ru_source"] for r in records)
    missing = [r for r in records if r["ru"] is None]
    extinct = Counter(r["extinct"] for r in records if r["extinct"])
    not_in_ioc = [r for r in records if sci_key(r["sci"]) not in ioc]
    ioc_only = sorted((v for k, v in ioc.items() if k not in avi_keys), key=lambda v: v["sci"])
    diffs = []
    for r in records:
        key = sci_key(r["sci"])
        a, b = (ioc.get(key) or {}).get("ru"), ebird.get(key)
        if a and b and a.casefold() != b.casefold():
            diffs.append((r["sci"], a, b))

    counts = {
        "total": total,
        "ru_present": total - len(missing),
        "ioc": by_source.get("ioc", 0),
        "ebird": by_source.get("ebird", 0),
        "wikidata": by_source.get("wikidata", 0),
        "manual": by_source.get("manual", 0),
        "missing": len(missing),
        "extinct": extinct.get("extinct", 0),
        "extinct_in_wild": extinct.get("extinct_in_wild", 0),
        "possibly_extinct": extinct.get("possibly_extinct", 0),
        "not_in_ioc": len(not_in_ioc),
        "ioc_not_in_avilist": len(ioc_only),
        "ioc_ebird_diff": len(diffs),
        "auth_leftover_initials": sum(1 for r in records
                                      if r["auth"] and LEFTOVER_INITIALS.search(r["auth"])),
    }
    dups = duplicate_ru(records)
    clashes = alt_clashes(records)
    heads = Counter(r["ru_head"] for r in records if r["ru_head"])
    counts.update({
        "ru_split_synonyms": len(names_log["split"]),
        "ru_with_alt": sum(1 for r in records if r["ru_alt"]),
        "ru_alt_total": sum(len(r["ru_alt"]) for r in records),
        "ru_paren_kept": len(names_log["paren_kept"]),
        "mixed_script_names": len(names_log["mixed"]),
        "ru_duplicates": len(dups),
        "ru_alt_clashes": len(clashes),
        "ru_heads": len(heads),
        "ru_head_manual": sum(1 for r in records if r["ru_head_source"] == "manual"),
        "ru_head_doubtful": len(names_log["head_doubtful"]),
        "ru_book_capital_kept": len(names_log["book_capital"]),
    })
    for rank in ("orders", "families"):
        src = Counter(t["ru_source"] for t in taxa[rank])
        counts[f"{rank}"] = len(taxa[rank])
        for s in ("zmmu", "wikidata", "manual"):
            counts[f"{rank}_ru_{s}"] = src.get(s, 0)
        counts[f"{rank}_ru_missing"] = src.get(None, 0)

    L = ["# Build report", ""]
    L += ["## Sources", ""]
    for name, state in status.items():
        L.append(f"- **{name}**: {state}")
    if warnings:
        L += ["", "### Override warnings", ""] + [f"- {w}" for w in warnings]
    L += ["", "## Counts", ""]
    L += md_table([
        ("Total species (AviList)", total),
        ("Russian name present", counts["ru_present"]),
        ("— from IOC", counts["ioc"]),
        ("— from eBird", counts["ebird"]),
        ("— from Wikidata", counts["wikidata"]),
        ("— manual override", counts["manual"]),
        ("Russian name missing", counts["missing"]),
        ("Extinct", counts["extinct"]),
        ("Extinct in the wild", counts["extinct_in_wild"]),
        ("Possibly extinct", counts["possibly_extinct"]),
        ("AviList species not in IOC (exact name)", counts["not_in_ioc"]),
        ("IOC species not in AviList (exact name)", counts["ioc_not_in_avilist"]),
        ("IOC vs eBird Russian name differs", counts["ioc_ebird_diff"]),
        ("`auth` with leftover initials", counts["auth_leftover_initials"]),
        ("Names with mixed Cyrillic/Latin words", counts["mixed_script_names"]),
        ("Russian names shared by several species", counts["ru_duplicates"]),
        ("`ru` split into `ru` + `ru_alt`", counts["ru_split_synonyms"]),
        ("Species with `ru_alt`", counts["ru_with_alt"]),
        ("`ru_alt` names in total", counts["ru_alt_total"]),
        ("`ru_alt` equal to another species' `ru`", counts["ru_alt_clashes"]),
        ("Distinct `ru_head`", counts["ru_heads"]),
        ("`ru_head` doubtful (auto)", counts["ru_head_doubtful"]),
        ("`ru_head` manual", counts["ru_head_manual"]),
        ("`ru_book` with the capital kept", counts["ru_book_capital_kept"]),
    ], ["Metric", "Count"])

    L += ["", "## Russian names: checks and index fields", ""]
    L += ["", f"### Words mixing Cyrillic and Latin letters ({len(names_log['mixed'])})", "",
          "Look-alike letters (e.g. Latin C in \"Cкоморох\"); fix in `data/overrides.csv`.", ""]
    L += md_table(names_log["mixed"], ["Scientific name", "Name"])
    L += ["", f"### Russian names shared by several species ({len(dups)})", "",
          "Not fixed automatically; resolve in `data/overrides.csv` when the right name is known.", ""]
    L += md_table([(name, "; ".join(f"*{r['sci']}* ({r['en']}, {r['ru_source']})" for r in recs))
                   for name, recs in dups], ["Russian name", "Species (English, source)"])
    L += ["", f"### Synonyms split out of `ru` ({len(names_log['split'])})", "",
          "Separators: comma, \" / \", \" или \", parentheses. The first part stays in `ru`, "
          "the rest go to `ru_alt` (first letter capitalised).", ""]
    L += md_table(names_log["split"], ["Scientific name", "Source name", "ru", "ru_alt", "ru_source"])
    L += ["", f"### Parentheses left in `ru` ({len(names_log['paren_kept'])})", "",
          "The parenthesis does not look like a name (comment, Latin, digits); not split.", ""]
    L += md_table(names_log["paren_kept"], ["Scientific name", "ru"])
    L += ["", f"### `ru_alt` equal to another species' `ru` ({len(clashes)})", ""]
    L += md_table(clashes, ["Scientific name", "ru_alt", "is `ru` of"])
    L += ["", "### Top 50 `ru_head` (index nests)", ""]
    L += md_table(heads.most_common(50), ["ru_head", "Species"])
    L += ["", f"### Doubtful `ru_head` ({len(names_log['head_doubtful'])})", "",
          "All-adjective names, a lower-case word after the head (apposition such as "
          "\"конёк пуны\" or a typo), or a capitalised head after other words. "
          "Fix systematic errors in `ru_head()`, single ones with `field = ru_head` in overrides.", ""]
    L += md_table(names_log["head_doubtful"], ["Scientific name", "ru", "ru_head"])
    L += ["", f"### `ru_book` with the capital kept ({len(names_log['book_capital'])})", "",
          "First word is a possessive adjective from a proper name.", ""]
    L += md_table([(sci, ru, en, "yes" if "'s" in (en or "") else "")
                   for sci, ru, en in names_log["book_capital"]],
                  ["Scientific name", "ru", "English", "English has 's"])

    L += ["", "## Orders and families: Russian names", "",
          "Sources: ZMMU passerine list family headers (`zmmu`), Wikidata (`wikidata`), "
          "`data/overrides.csv` (`manual`). Written to `data/taxa.json`.", ""]
    L += md_table([(rank, counts[rank], counts[f"{rank}_ru_zmmu"], counts[f"{rank}_ru_wikidata"],
                    counts[f"{rank}_ru_manual"], counts[f"{rank}_ru_missing"])
                   for rank in ("orders", "families")],
                  ["Rank", "Total", "zmmu", "wikidata", "manual", "missing"])
    if taxa_warnings:
        L += ["", "### Taxon name warnings", ""] + [f"- {w}" for w in taxa_warnings]
    missing_taxa = [("order", t["sci"], "", t["count"]) for t in taxa["orders"] if not t["ru"]]
    missing_taxa += [("family", t["sci"], t["en"], t["count"]) for t in taxa["families"]
                     if not t["ru"]]
    L += ["", f"### Orders and families without a Russian name ({len(missing_taxa)})", ""]
    L += md_table(missing_taxa, ["Rank", "Scientific name", "English", "Species"])

    L += ["", f"## AviList species not found in IOC ({len(not_in_ioc)})", ""]
    L += md_table([(r["sci"], r["en"], r["ru_source"] or "—") for r in not_in_ioc],
                  ["Scientific name", "English", "ru_source"])
    L += ["", f"## Species with no Russian name anywhere ({len(missing)})", ""]
    L += md_table([(r["sci"], r["en"]) for r in missing], ["Scientific name", "English"])
    L += ["", f"## IOC species not found in AviList ({len(ioc_only)})", ""]
    L += md_table([(v["sci"], v["en"], v["ru"]) for v in ioc_only],
                  ["IOC scientific name", "English (IOC)", "Russian (IOC)"])
    L += ["", f"## IOC and eBird Russian names differ ({len(diffs)})", "",
          "Report only; the JSON always prefers IOC.", ""]
    L += md_table(diffs, ["Scientific name", "IOC", "eBird"])
    OUT_REPORT.write_text("\n".join(L) + "\n", encoding="utf-8")
    return counts


# --------------------------------------------------------------------------- main

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    g = ap.add_mutually_exclusive_group()
    g.add_argument("--offline", action="store_true", help="use only files already on disk")
    g.add_argument("--fetch", action="store_true", help="force re-download of all sources")
    args = ap.parse_args()
    mode = "offline" if args.offline else "fetch" if args.fetch else "default"
    log(f"mode: {mode}")
    _test_normalise_authority()
    _test_names()

    status = {}
    ensure_file(AVILIST_PATH, AVILIST_URL, mode, required=True)
    ensure_file(IOC_PATH, IOC_URL, mode, required=True)
    status["AviList v2025b"] = f"`{AVILIST_PATH.relative_to(ROOT)}`"
    status["IOC Multilingual 15.2"] = f"`{IOC_PATH.relative_to(ROOT)}`"

    ebird = {}
    if ensure_file(EBIRD_PATH, EBIRD_URL, mode, required=False):
        ebird = load_ebird(EBIRD_PATH)
        status["eBird taxonomy (ru)"] = f"`{EBIRD_PATH.relative_to(ROOT)}`, {len(ebird)} species with Cyrillic names"
    else:
        status["eBird taxonomy (ru)"] = "SKIPPED (not available)"

    wikidata = {}
    if mode == "fetch" or (mode == "default" and not WIKIDATA_PATH.exists()):
        try:
            fetch_wikidata(WIKIDATA_PATH)
        except Exception as exc:  # noqa: BLE001
            log(f"  WARNING: Wikidata query failed: {exc}")
    if WIKIDATA_PATH.exists():
        wikidata = load_wikidata(WIKIDATA_PATH)
        status["Wikidata P1843 (ru)"] = f"`{WIKIDATA_PATH.relative_to(ROOT)}`, {len(wikidata)} taxa"
    else:
        status["Wikidata P1843 (ru)"] = "SKIPPED (unreachable or no cache)"

    log("loading AviList and IOC")
    records = load_avilist(AVILIST_PATH)
    ioc = load_ioc(IOC_PATH)

    apply_cascade(records, ioc, ebird, wikidata)
    overrides = load_overrides(OVERRIDES_PATH)
    taxa = build_taxa(records)
    taxon_keys = {sci_key(t["sci"]) for t in taxa["orders"] + taxa["families"]}
    species_keys = {sci_key(r["sci"]) for r in records}
    warnings = apply_overrides(records, overrides, taxon_keys)
    names_log = normalise_names(records)
    finalise(records)

    # order / family Russian names
    taxa_warnings: list[str] = []
    zmmu = {}
    if ensure_file(ZMMU_PASSERINE_PATH, ZMMU_PASSERINE_URL, mode, required=False):
        zmmu, taxa_warnings = load_zmmu_families(
            ZMMU_PASSERINE_PATH, {sci_key(r["sci"]): r["family"] for r in records},
            {t["sci"] for t in taxa["families"]})
        status["ZMMU passerine list (family names)"] = (
            f"`{ZMMU_PASSERINE_PATH.relative_to(ROOT)}`, {len(zmmu)} families")
    else:
        status["ZMMU passerine list (family names)"] = "SKIPPED (not available)"
    taxon_names = sorted(t["sci"] for t in taxa["orders"] + taxa["families"])
    cached_names = None
    if WIKIDATA_TAXA_PATH.exists():
        try:
            cached_names = json.loads(WIKIDATA_TAXA_PATH.read_text(encoding="utf-8"))["queried"]
        except Exception:  # noqa: BLE001 - corrupt cache: refetch
            cached_names = None
    if mode == "fetch" or (mode == "default" and (
            cached_names is None or not set(taxon_names) <= set(cached_names))):
        try:
            fetch_wikidata_taxa(WIKIDATA_TAXA_PATH, taxon_names)
        except Exception as exc:  # noqa: BLE001
            log(f"  WARNING: Wikidata order/family query failed: {exc}")
    wikidata_taxa = {}
    if WIKIDATA_TAXA_PATH.exists():
        wikidata_taxa = load_wikidata_taxa(WIKIDATA_TAXA_PATH)
        status["Wikidata orders/families (ru)"] = (
            f"`{WIKIDATA_TAXA_PATH.relative_to(ROOT)}`, {len(wikidata_taxa)} taxa")
    else:
        status["Wikidata orders/families (ru)"] = "SKIPPED (unreachable or no cache)"
    apply_taxa_names(taxa, zmmu, wikidata_taxa)
    taxa_warnings += apply_taxa_overrides(
        taxa, [ov for ov in overrides
               if sci_key(ov["sci"]) in taxon_keys and sci_key(ov["sci"]) not in species_keys])

    write_json(records, taxa)
    counts = write_report(records, ioc, ebird, wikidata, status, warnings, taxa, taxa_warnings,
                          names_log)

    print(f"Wrote {OUT_JSON.relative_to(ROOT)}, {OUT_TAXA_JSON.relative_to(ROOT)}, "
          f"{OUT_DOCS_JSON.relative_to(ROOT)}, {OUT_DOCS_TAXA_JSON.relative_to(ROOT)}, "
          f"{OUT_REPORT.relative_to(ROOT)}")
    for name, state in status.items():
        print(f"  {name}: {state}")
    for k, v in counts.items():
        print(f"  {k}: {v}")
    for w in taxa_warnings:
        print(f"  taxon warning: {w}")
    for rank in ("orders", "families"):
        missing = [t["sci"] for t in taxa[rank] if not t["ru"]]
        print(f"  {rank} without Russian name ({len(missing)}): {', '.join(missing) or '-'}")


if __name__ == "__main__":
    main()
